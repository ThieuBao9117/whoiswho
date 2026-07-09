import os
import sys
from datetime import datetime, timezone
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from sqlalchemy.orm import Session
from sqlalchemy import extract
from app.core.database import SessionLocal
from app.models.csb_employee_ref import CSBEmployeeRef
from app.models.crw_models import CRWConnection, ConnectionStatus
from app.services.game_logic import get_leaderboard_data

def get_dept(emp):
    if not emp: return "N/A"
    return (emp.part or emp.department) or "N/A"

def run_export():
    db: Session = SessionLocal()
    
    # We will export from April 2026 to current month
    start_month = 4
    start_year = 2026
    current_date = datetime.now()
    end_month = current_date.month
    end_year = current_date.year
    
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)
    
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    winner_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

    has_data = False
    
    for year in range(start_year, end_year + 1):
        m_start = start_month if year == start_year else 1
        m_end = end_month if year == end_year else 12
        
        for month in range(m_start, m_end + 1):
            winners = get_leaderboard_data(db, limit=1000, target_month=month, target_year=year)
            if not winners: continue
            
            for winner in winners:
                has_data = True
                emp_id = winner["id"]
                emp = db.query(CSBEmployeeRef).filter(CSBEmployeeRef.id == emp_id).first()
                if not emp: continue
                
                # Create a sheet for this specific winner
                sheet_title = f"{winner['emp_code']}_T{month}"
                ws = wb.create_sheet(title=sheet_title[:31]) # Max length is 31
                
                # Write winner header
                ws.append([f"CHI TIẾT KẾT NỐI - NGƯỜI THẮNG CUỘC THÁNG {month}/{year}"])
                ws.append([])
                ws.append([
                    f"Người thắng: {winner['name']} ({winner['emp_code']})", 
                    f"Phòng ban: {winner['dept']}", 
                    f"Tổng số lượng kết nối: {winner['raw_count']} người",
                    f"Thời gian hoàn thành KPI: {winner['score']}"
                ])
                for cell in ws._cells.values():
                    if cell.row == ws.max_row:
                        cell.fill = winner_fill
                        cell.font = Font(bold=True)
                
                ws.append([]) # Empty row
                
                # Table headers for connections
                ws.append(["STT", "Người được kết nối (Bạn bè)", "Mã NV", "Phòng ban / Tổ", "Chức vụ", "Thời gian kết nối thành công"])
                for cell in ws._cells.values():
                    if cell.row == ws.max_row:
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = Alignment(horizontal="center")
                        
                from sqlalchemy import or_
                conns = db.query(CRWConnection).filter(
                    or_(CRWConnection.connector_id == emp_id, CRWConnection.new_hire_id == emp_id),
                    CRWConnection.status == ConnectionStatus.ACCEPTED
                ).all()
                conns.sort(key=lambda x: x.responded_at or x.created_at)
                
                for idx, conn in enumerate(conns, 1):
                    # Determine who is the "friend" (the other person)
                    if conn.connector_id == emp_id:
                        target = conn.new_hire
                    else:
                        target = conn.connector
                        
                    t_name = target.full_name if target else "Unknown"
                    t_code = target.emp_code if target else "Unknown"
                    t_dept = get_dept(target)
                    t_role = target.role if target else "N/A"
                    c_time = (conn.responded_at or conn.created_at)
                    c_time_str = c_time.strftime("%Y-%m-%d %H:%M:%S") if c_time else "N/A"
                    
                    ws.append([idx, t_name, t_code, t_dept, t_role, c_time_str])
                    
                # Adjust widths
                for col in ws.columns:
                    max_length = 0
                    column = col[0].column_letter
                    for cell in col:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(cell.value)
                        except:
                            pass
                    ws.column_dimensions[column].width = min(max_length + 2, 50)
                
    if not has_data:
        ws = wb.create_sheet(title="No Data")
        ws.append(["Chưa có người chiến thắng nào trong hệ thống."])
        
    output_path = os.path.join(os.path.dirname(__file__), "Chi_tiet_nguoi_chien_thang_theo_Sheet_v2.xlsx")
    wb.save(output_path)
    print(f"Exported successfully to {output_path}")

if __name__ == "__main__":
    run_export()
